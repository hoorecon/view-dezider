from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.space_before = Pt(2)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return h

def add_table_row(table, cells_data, bold=False, header=False):
    row = table.add_row()
    for i, text in enumerate(cells_data):
        cell = row.cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(str(text))
        run.font.size = Pt(9)
        run.font.name = 'Calibri'
        if bold or header:
            run.bold = True
        if header:
            from lxml import etree
            shading = etree.SubElement(cell._tc.get_or_add_tcPr(), qn('w:shd'))
            shading.set(qn('w:fill'), '1F2937')
            shading.set(qn('w:val'), 'clear')
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ===== TITLE =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('DATA SECURITY DECLARATION')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('DigiLocker / API Setu Partner Integration')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

# Meta info
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_run = meta.add_run('Organization: VEALES (View Dezider Platform)  |  Date: 02 May 2026  |  Version: 1.0  |  Classification: Confidential')
meta_run.font.size = Pt(8)
meta_run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

doc.add_paragraph('─' * 80)

# ===== 1. PURPOSE =====
add_heading_styled('1. Purpose', level=2)
doc.add_paragraph(
    'This declaration outlines the data security practices, encryption standards, user consent mechanisms, '
    'and data handling policies implemented by VEALES ("the Company") for the View Dezider platform in relation '
    'to the integration with DigiLocker / API Setu services for identity verification (eKYC) of platform participants.'
)

# ===== 2. DATA COLLECTED =====
add_heading_styled('2. Data Collected via DigiLocker', level=2)
doc.add_paragraph(
    'The following data points are fetched only after explicit user consent via the DigiLocker OAuth2 authorization flow. '
    'No Aadhaar number is stored at any point.'
)

t2 = doc.add_table(rows=1, cols=3)
t2.style = 'Table Grid'
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_row(t2, ['Data Point', 'Source Document', 'Purpose'], header=True)
# Remove first empty row
t2.rows[0]._tr.getparent().remove(t2.rows[0]._tr)
add_table_row(t2, ['Full Name', 'Aadhaar', 'Verify participant identity in group decisions'])
add_table_row(t2, ['Date of Birth', 'Aadhaar', 'Age verification for compliance'])
add_table_row(t2, ['Gender', 'Aadhaar', 'Optional demographic for SME profiling'])
add_table_row(t2, ['Photo', 'Aadhaar', 'Visual identity confirmation by session admin'])
add_table_row(t2, ['Udyam / CIN Number', 'Udyam Certificate', 'Verify business entity legitimacy'])

# ===== 3. ENCRYPTION =====
add_heading_styled('3. Encryption Standards', level=2)

t3 = doc.add_table(rows=1, cols=3)
t3.style = 'Table Grid'
t3.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_row(t3, ['Layer', 'Standard', 'Implementation'], header=True)
t3.rows[0]._tr.getparent().remove(t3.rows[0]._tr)
add_table_row(t3, ['Data in Transit', 'TLS 1.3', 'All API calls encrypted via HTTPS with TLS 1.3'])
add_table_row(t3, ['Data at Rest', 'AES-256', 'KYC records in MongoDB encrypted at rest using AES-256'])
add_table_row(t3, ['Token Storage', 'HMAC-SHA256', 'OAuth2 tokens hashed before storage; plaintext never persisted'])
add_table_row(t3, ['Password Hashing', 'bcrypt (cost 12)', 'User passwords hashed using bcrypt, never stored in plaintext'])
add_table_row(t3, ['API Authentication', 'JWT (RS256)', 'All internal API calls authenticated via signed JWTs'])

# ===== 4. CONSENT FLOW =====
add_heading_styled('4. User Consent Flow', level=2)
doc.add_paragraph('The following multi-step consent mechanism is enforced before any DigiLocker data is accessed:')

steps = [
    ('Informed Consent Screen', 'User is shown a clear explanation of what data will be fetched, why it is needed, and how it will be used.'),
    ('DigiLocker OAuth Authorization', 'User is redirected to the official DigiLocker portal where they independently authenticate and grant consent.'),
    ('Granular Document Selection', 'User selects which documents to share within the DigiLocker interface. No documents are fetched without user selection.'),
    ('Post-Verification Confirmation', 'After verification, user is shown exactly what data was retrieved and given the option to revoke access.'),
    ('Revocation Right', 'Users can request complete deletion of their KYC data at any time via app settings or by contacting support.'),
]
for i, (title_text, desc) in enumerate(steps, 1):
    p = doc.add_paragraph()
    run_num = p.add_run(f'Step {i}: {title_text} — ')
    run_num.bold = True
    run_num.font.size = Pt(10)
    run_desc = p.add_run(desc)
    run_desc.font.size = Pt(10)

# ===== 5. NO THIRD-PARTY SHARING =====
add_heading_styled('5. No Third-Party Data Sharing', level=2)
doc.add_paragraph('We hereby declare that:')

declarations = [
    'NO personally identifiable information (PII) obtained via DigiLocker is shared with any third party, partner, advertiser, analytics provider, or external service.',
    'NO Aadhaar number, biometric data, or raw document images are stored on our servers.',
    'NO KYC data is used for any purpose other than participant identity verification within the View Dezider platform.',
    'NO data is transferred outside the territory of India. All servers and databases are hosted within India.',
    'NO automated profiling, scoring, or AI training is performed on KYC data.',
]
for d in declarations:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(d)
    run.font.size = Pt(10)

# ===== 6. DATA RETENTION =====
add_heading_styled('6. Data Retention & Deletion Policy', level=2)

t6 = doc.add_table(rows=1, cols=3)
t6.style = 'Table Grid'
t6.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_row(t6, ['Scenario', 'Retention Period', 'Action'], header=True)
t6.rows[0]._tr.getparent().remove(t6.rows[0]._tr)
add_table_row(t6, ['Active verified user', 'Duration of active account', 'Data retained for verification status only'])
add_table_row(t6, ['User requests deletion', 'Within 48 hours', 'All KYC records permanently deleted'])
add_table_row(t6, ['Account inactive > 12 months', 'Auto-purge', 'KYC data auto-deleted after inactivity'])
add_table_row(t6, ['DigiLocker OAuth tokens', '24 hours max', 'Tokens auto-expire and purged from storage'])

# ===== 7. INCIDENT RESPONSE =====
add_heading_styled('7. Incident Response', level=2)
doc.add_paragraph('In the event of any data breach or security incident involving DigiLocker-sourced data:')
items_7 = [
    'CERT-In will be notified within 6 hours as per Indian IT Act requirements.',
    'Affected users will be notified within 24 hours with details of the breach and remediation steps.',
    'DigiLocker / API Setu team will be informed immediately for coordinated response.',
]
for item in items_7:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(item)
    run.font.size = Pt(10)

# ===== 8. COMPLIANCE =====
add_heading_styled('8. Regulatory Compliance', level=2)
doc.add_paragraph('This implementation complies with:')
compliances = [
    'Information Technology Act, 2000 (India)',
    'IT (Reasonable Security Practices and Procedures) Rules, 2011',
    'UIDAI Aadhaar Data Vault Guidelines',
    'Digital Personal Data Protection Act (DPDPA), 2023',
    'API Setu Terms of Use and Privacy Statement',
]
for c in compliances:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(c)
    run.font.size = Pt(10)

# ===== 9. SIGNATORY =====
doc.add_paragraph('─' * 80)
add_heading_styled('9. Authorized Signatory', level=2)

sig_table = doc.add_table(rows=6, cols=2)
sig_table.style = 'Table Grid'
sig_data = [
    ('Name', 'A D Shezhiyan Raj'),
    ('Designation', 'Founder & CEO'),
    ('Organization', 'VEALES'),
    ('Platform', 'View Dezider'),
    ('Email', '[Your Email]'),
    ('Contact', '[Your Phone]'),
]
for i, (label, value) in enumerate(sig_data):
    sig_table.rows[i].cells[0].text = label
    sig_table.rows[i].cells[1].text = value
    for cell in sig_table.rows[i].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)

doc.add_paragraph()
p_sig = doc.add_paragraph()
p_sig.add_run('\nSignature: ').bold = True
p_sig.add_run('____________________________')

p_date = doc.add_paragraph()
p_date.add_run('\nDate: ').bold = True
p_date.add_run('02 May 2026')

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('This document is issued by VEALES in support of the DigiLocker / API Setu partner registration application.\nThe declarations made herein are accurate and binding.')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
run.italic = True

# Save
output_path = '/app/backend/static/ViewDezider_Data_Security_Declaration.docx'
os.makedirs('/app/backend/static', exist_ok=True)
doc.save(output_path)
print(f"DOCX saved to {output_path}")
