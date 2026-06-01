"""
Solution Store — bulk factor-value ingestion.

Three input methods, all funnelling into an admin-reviewed submission queue:
  1. XLS template  : download template -> fill -> upload (multipart).
  2. Google Sheet  : a "Published to web" / shareable sheet link, read as CSV.
  3. Webhook API   : per-solution secret token; 3rd parties POST JSON for live updates.

Every submission lands in `store_factor_submissions` with status='pending' and must be
approved by an admin before the values are written onto the solution's
`quantitative_factors`. (A per-token `auto_approve` flag — admin-controlled — can apply
webhook pushes immediately.)
"""
import io
import csv
import re
import uuid
import secrets
import logging
from datetime import datetime, timezone

import httpx
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Header, Query
from fastapi.responses import StreamingResponse

from core.auth import get_current_user
from core.database import db

logger = logging.getLogger("routes.store_ingestion")
router = APIRouter()

ADMIN_ROLES = {"super_admin", "co_admin", "admin"}
TEMPLATE_COLUMNS = [
    "solution_id", "factor_name", "factor_type", "value",
    "unit", "currency", "source_note", "effective_date",
]
WEBHOOK_PATH = "/api/solutions-store/factor-values/webhook"


def _now():
    return datetime.now(timezone.utc).isoformat()


def _is_admin(user: dict) -> bool:
    return user.get("role") in ADMIN_ROLES


def _norm(s) -> str:
    return str(s or "").strip().lower()


def _coerce_value(value, factor_type: str):
    """Quantitative -> number when possible; qualitative -> trimmed string."""
    if value is None:
        return None
    if _norm(factor_type) == "quantitative":
        try:
            num = float(value)
            return int(num) if num.is_integer() else num
        except (TypeError, ValueError):
            return str(value).strip()
    return str(value).strip()


def _row_from_record(rec: dict, default_solution_id=None) -> dict | None:
    """Normalise one input record (dict keyed by template columns) into a factor row."""
    def get(k):
        if not rec:
            return None
        v = rec.get(k)
        return v if v is not None else rec.get(k.upper())
    sid = (get("solution_id") or default_solution_id or "").strip() if (get("solution_id") or default_solution_id) else ""
    fname = str(get("factor_name") or "").strip()
    if not sid or not fname:
        return None
    ftype = _norm(get("factor_type")) or "quantitative"
    if ftype not in ("quantitative", "qualitative"):
        ftype = "quantitative"
    return {
        "solution_id": sid,
        "factor_name": fname,
        "factor_type": ftype,
        "value": _coerce_value(get("value"), ftype),
        "unit": (str(get("unit")).strip() if get("unit") not in (None, "") else None),
        "currency": (str(get("currency")).strip() if get("currency") not in (None, "") else None),
        "source_note": (str(get("source_note")).strip() if get("source_note") not in (None, "") else None),
        "effective_date": (str(get("effective_date")).strip() if get("effective_date") not in (None, "") else None),
    }


# ──────────────────────────────────────────────────────────────────────────
# 1) XLS template download
# ──────────────────────────────────────────────────────────────────────────
@router.get("/solutions-store/factor-template.xlsx")
async def download_factor_template(
    solution_id: str = Query(None, description="Optional: prefill with this solution's current factors"),
    user: dict = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Factor Values"

    header_fill = PatternFill(start_color="8E24AA", end_color="8E24AA", fill_type="solid")
    bold_white = Font(bold=True, color="FFFFFF")
    for ci, col in enumerate(TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=ci, value=col)
        c.font = bold_white
        c.fill = header_fill
    widths = [38, 26, 14, 12, 10, 10, 30, 16]
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    prefilled = False
    if solution_id:
        sol = await db.solutions_store.find_one({"solution_id": solution_id}, {"_id": 0})
        if sol and (sol.get("created_by") == user["user_id"] or _is_admin(user)):
            r = 2
            for qf in (sol.get("quantitative_factors") or []):
                ws.cell(row=r, column=1, value=solution_id)
                ws.cell(row=r, column=2, value=qf.get("factor_name"))
                ws.cell(row=r, column=3, value="qualitative" if qf.get("data_type") == "text" else "quantitative")
                ws.cell(row=r, column=4, value=qf.get("value"))
                ws.cell(row=r, column=5, value=qf.get("unit"))
                ws.cell(row=r, column=6, value=qf.get("currency"))
                ws.cell(row=r, column=7, value=qf.get("source_note"))
                r += 1
            prefilled = r > 2
    if not prefilled:
        ws.cell(row=2, column=1, value=solution_id or "<your-solution-id>")
        ws.cell(row=2, column=2, value="Monthly Cost")
        ws.cell(row=2, column=3, value="quantitative")
        ws.cell(row=2, column=4, value=2500)
        ws.cell(row=2, column=5, value="INR")
        ws.cell(row=2, column=6, value="INR")
        ws.cell(row=2, column=7, value="From official pricing page")

    info = wb.create_sheet("Instructions")
    notes = [
        "How to use this template",
        "",
        "1. One row per factor value. Repeat solution_id across rows for the same solution.",
        "2. factor_type must be 'quantitative' (a number) or 'qualitative' (text).",
        "3. value: a number for quantitative factors; short text for qualitative.",
        "4. unit / currency / source_note are optional. effective_date is optional (YYYY-MM-DD).",
        "5. factor_name should match the factor you want to populate.",
        "6. Save as .xlsx and upload it back in the app (Bulk Factor Updates).",
        "7. All uploads are reviewed & approved by an admin before going live.",
    ]
    for i, line in enumerate(notes, start=1):
        cell = info.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=factor_values_template.xlsx"},
    )


# ──────────────────────────────────────────────────────────────────────────
# Submission helpers
# ──────────────────────────────────────────────────────────────────────────
async def _create_submissions(rows: list, user_id: str, user_name: str, source: str,
                              allow_solutions: set | None = None):
    """Group rows by solution_id, validate ownership, create pending submissions.
    `allow_solutions` (if given) restricts which solution_ids are permitted (webhook token)."""
    by_sid = {}
    for r in rows:
        by_sid.setdefault(r["solution_id"], []).append(r)

    created, errors = [], []
    for sid, srows in by_sid.items():
        if allow_solutions is not None and sid not in allow_solutions:
            errors.append({"solution_id": sid, "error": "Not permitted for this token"})
            continue
        sol = await db.solutions_store.find_one({"solution_id": sid}, {"_id": 0, "name": 1, "created_by": 1})
        if not sol:
            errors.append({"solution_id": sid, "error": "Solution not found"})
            continue
        if allow_solutions is None and sol.get("created_by") != user_id and user_id != "__admin__":
            # ownership/admin enforced by caller passing user_id=="__admin__" for admins
            errors.append({"solution_id": sid, "error": "Not your solution"})
            continue
        sub = {
            "submission_id": str(uuid.uuid4()),
            "solution_id": sid,
            "solution_name": sol.get("name", ""),
            "submitted_by": user_id,
            "submitted_by_name": user_name,
            "source": source,
            "status": "pending",
            "rows": srows,
            "row_count": len(srows),
            "created_at": _now(),
            "reviewed_by": None,
            "reviewed_at": None,
            "note": None,
        }
        await db.store_factor_submissions.insert_one(sub)
        created.append({"submission_id": sub["submission_id"], "solution_id": sid, "rows": len(srows)})
    return created, errors


async def _apply_submission(sub: dict):
    """Upsert the submission rows onto the solution's quantitative_factors."""
    sol = await db.solutions_store.find_one({"solution_id": sub["solution_id"]}, {"_id": 0, "quantitative_factors": 1})
    if not sol:
        raise HTTPException(status_code=404, detail="Solution no longer exists")
    qfs = list(sol.get("quantitative_factors") or [])
    index = {_norm(qf.get("factor_name")): i for i, qf in enumerate(qfs)}
    for row in sub["rows"]:
        entry = {
            "factor_name": row["factor_name"],
            "value": row["value"],
            "unit": row.get("unit"),
            "data_type": "text" if row.get("factor_type") == "qualitative" else "numeric",
            "currency": row.get("currency"),
            "source_note": row.get("source_note"),
            "effective_date": row.get("effective_date"),
            "updated_at": _now(),
            "updated_via": sub["source"],
        }
        key = _norm(row["factor_name"])
        if key in index:
            qfs[index[key]] = {**qfs[index[key]], **entry}
        else:
            index[key] = len(qfs)
            qfs.append(entry)
    await db.solutions_store.update_one(
        {"solution_id": sub["solution_id"]},
        {"$set": {"quantitative_factors": qfs, "updated_at": _now()}},
    )


# ──────────────────────────────────────────────────────────────────────────
# 2) XLS upload
# ──────────────────────────────────────────────────────────────────────────
@router.post("/solutions-store/factor-values/upload")
async def upload_factor_xls(
    file: UploadFile = File(...),
    solution_id: str = Query(None, description="Optional default solution_id for rows missing one"),
    user: dict = Depends(get_current_user),
):
    from openpyxl import load_workbook

    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx file")
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 5MB)")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb["Factor Values"] if "Factor Values" in wb.sheetnames else wb.worksheets[0]
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read spreadsheet: {str(e)[:120]}")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [(_norm(h)) for h in next(rows_iter)]
    except StopIteration:
        raise HTTPException(status_code=400, detail="Spreadsheet is empty")
    col_idx = {c: header.index(c) for c in TEMPLATE_COLUMNS if c in header}
    if "factor_name" not in col_idx or "value" not in col_idx:
        raise HTTPException(status_code=400, detail="Template must include at least 'factor_name' and 'value' columns")

    records = []
    for raw in rows_iter:
        if raw is None or all(v in (None, "") for v in raw):
            continue
        rec = {c: (raw[i] if i < len(raw) else None) for c, i in col_idx.items()}
        records.append(rec)

    parsed = [r for r in (_row_from_record(rec, solution_id) for rec in records) if r]
    if not parsed:
        raise HTTPException(status_code=400, detail="No valid rows found (need solution_id + factor_name)")

    uid = "__admin__" if _is_admin(user) else user["user_id"]
    created, errors = await _create_submissions(parsed, uid, user.get("name", ""), "xls")
    if not created and errors:
        raise HTTPException(status_code=400, detail=errors[0]["error"])
    return {"submissions": created, "errors": errors, "parsed_rows": len(parsed), "status": "pending_review"}


# ──────────────────────────────────────────────────────────────────────────
# 3) Google Sheet (published CSV) import
# ──────────────────────────────────────────────────────────────────────────
def _gsheet_to_csv_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        raise HTTPException(status_code=400, detail="Provide a Google Sheet link")
    if "output=csv" in url or "format=csv" in url:
        return url
    if "/pub" in url:  # Published-to-web link
        sep = "&" if "?" in url else "?"
        return f"{url}{sep}output=csv"
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url)
    if m:
        sheet_id = m.group(1)
        gid_m = re.search(r"[#&?]gid=([0-9]+)", url)
        gid = gid_m.group(1) if gid_m else "0"
        return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
    raise HTTPException(status_code=400, detail="Unrecognised Google Sheet link. Use 'Publish to web' (CSV) or a share link.")


@router.post("/solutions-store/factor-values/import-gsheet")
async def import_factor_gsheet(request: Request, user: dict = Depends(get_current_user)):
    body = await request.json()
    csv_url = _gsheet_to_csv_url(body.get("sheet_url"))
    default_sid = (body.get("solution_id") or "").strip() or None
    try:
        async with httpx.AsyncClient(follow_redirects=True, timeout=20) as client:
            resp = await client.get(csv_url)
        if resp.status_code != 200 or "html" in resp.headers.get("content-type", "").lower():
            raise HTTPException(status_code=400, detail="Could not read the sheet as CSV. Ensure it is 'Published to web' or link-shareable (Anyone with the link).")
        text = resp.text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to fetch sheet: {str(e)[:120]}")

    reader = csv.DictReader(io.StringIO(text))
    records = [{(_norm(k)): v for k, v in row.items()} for row in reader]
    parsed = [r for r in (_row_from_record(rec, default_sid) for rec in records) if r]
    if not parsed:
        raise HTTPException(status_code=400, detail="No valid rows found. First row must be headers including solution_id, factor_name, value.")

    uid = "__admin__" if _is_admin(user) else user["user_id"]
    created, errors = await _create_submissions(parsed, uid, user.get("name", ""), "gsheet")
    if not created and errors:
        raise HTTPException(status_code=400, detail=errors[0]["error"])
    return {"submissions": created, "errors": errors, "parsed_rows": len(parsed), "status": "pending_review"}


# ──────────────────────────────────────────────────────────────────────────
# 4) Per-solution ingestion token (for the 3rd-party Webhook)
# ──────────────────────────────────────────────────────────────────────────
def _sample_payload(solution_id: str) -> dict:
    return {
        "rows": [
            {"factor_name": "Monthly Cost", "factor_type": "quantitative", "value": 2499, "unit": "INR", "currency": "INR", "source_note": "live price feed"},
            {"factor_name": "Availability", "factor_type": "qualitative", "value": "In stock"},
        ]
    }


async def _require_solution_owner(solution_id: str, user: dict):
    sol = await db.solutions_store.find_one({"solution_id": solution_id}, {"_id": 0, "created_by": 1, "name": 1})
    if not sol:
        raise HTTPException(status_code=404, detail="Solution not found")
    if sol.get("created_by") != user["user_id"] and not _is_admin(user):
        raise HTTPException(status_code=403, detail="Not your solution")
    return sol


@router.get("/solutions-store/{solution_id}/ingestion-token")
async def get_ingestion_token(solution_id: str, user: dict = Depends(get_current_user)):
    await _require_solution_owner(solution_id, user)
    tok = await db.store_ingestion_tokens.find_one({"solution_id": solution_id}, {"_id": 0})
    return {
        "solution_id": solution_id,
        "token": tok.get("token") if tok else None,
        "auto_approve": bool(tok.get("auto_approve")) if tok else False,
        "webhook_path": WEBHOOK_PATH,
        "header_name": "X-Ingestion-Token",
        "sample_payload": _sample_payload(solution_id),
    }


@router.post("/solutions-store/{solution_id}/ingestion-token")
async def generate_ingestion_token(solution_id: str, user: dict = Depends(get_current_user)):
    await _require_solution_owner(solution_id, user)
    token = "sk_ing_" + secrets.token_urlsafe(24)
    existing = await db.store_ingestion_tokens.find_one({"solution_id": solution_id}, {"_id": 0, "auto_approve": 1})
    await db.store_ingestion_tokens.update_one(
        {"solution_id": solution_id},
        {"$set": {
            "solution_id": solution_id,
            "token": token,
            "created_by": user["user_id"],
            "created_at": _now(),
            "auto_approve": bool(existing.get("auto_approve")) if existing else False,
        }},
        upsert=True,
    )
    return {
        "solution_id": solution_id,
        "token": token,
        "webhook_path": WEBHOOK_PATH,
        "header_name": "X-Ingestion-Token",
        "sample_payload": _sample_payload(solution_id),
    }


@router.post("/solutions-store/factor-values/webhook")
async def factor_values_webhook(request: Request, x_ingestion_token: str = Header(None)):
    """3rd-party live updates. Auth via X-Ingestion-Token header (no user session).
    Creates a pending submission (or applies immediately if the token is auto_approve)."""
    if not x_ingestion_token:
        raise HTTPException(status_code=401, detail="Missing X-Ingestion-Token header")
    tok = await db.store_ingestion_tokens.find_one({"token": x_ingestion_token}, {"_id": 0})
    if not tok:
        raise HTTPException(status_code=401, detail="Invalid ingestion token")
    solution_id = tok["solution_id"]

    body = await request.json()
    raw_rows = body.get("rows") or body.get("factors") or []
    if not isinstance(raw_rows, list) or not raw_rows:
        raise HTTPException(status_code=400, detail="Body must include a non-empty 'rows' array")
    parsed = [r for r in (_row_from_record({**(rec or {}), "solution_id": solution_id}, solution_id) for rec in raw_rows) if r]
    if not parsed:
        raise HTTPException(status_code=400, detail="No valid rows (each needs factor_name + value)")

    created, errors = await _create_submissions(
        parsed, tok.get("created_by", "webhook"), "Webhook", "webhook", allow_solutions={solution_id}
    )
    if not created:
        raise HTTPException(status_code=400, detail=(errors[0]["error"] if errors else "Nothing ingested"))

    applied = False
    if tok.get("auto_approve"):
        for c in created:
            sub = await db.store_factor_submissions.find_one({"submission_id": c["submission_id"]})
            await _apply_submission(sub)
            await db.store_factor_submissions.update_one(
                {"submission_id": c["submission_id"]},
                {"$set": {"status": "approved", "reviewed_by": "auto", "reviewed_at": _now()}},
            )
        applied = True
    return {
        "status": "applied" if applied else "pending_review",
        "submissions": created,
        "rows": len(parsed),
    }


# ──────────────────────────────────────────────────────────────────────────
# 5) Admin review queue
# ──────────────────────────────────────────────────────────────────────────
@router.get("/solutions-store/factor-submissions")
async def list_factor_submissions(
    status: str = Query("pending"),
    user: dict = Depends(get_current_user),
):
    q = {}
    if status and status != "all":
        q["status"] = status
    if not _is_admin(user):
        q["submitted_by"] = user["user_id"]  # owners see only their own
    items = await db.store_factor_submissions.find(q, {"_id": 0}).sort("created_at", -1).to_list(200)
    return {"items": items, "count": len(items)}


@router.post("/solutions-store/factor-submissions/{submission_id}/approve")
async def approve_factor_submission(submission_id: str, user: dict = Depends(get_current_user)):
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="Admin approval required")
    sub = await db.store_factor_submissions.find_one({"submission_id": submission_id})
    if not sub:
        raise HTTPException(status_code=404, detail="Submission not found")
    if sub["status"] != "pending":
        raise HTTPException(status_code=400, detail=f"Already {sub['status']}")
    await _apply_submission(sub)
    await db.store_factor_submissions.update_one(
        {"submission_id": submission_id},
        {"$set": {"status": "approved", "reviewed_by": user["user_id"], "reviewed_at": _now()}},
    )
    return {"status": "approved", "solution_id": sub["solution_id"], "applied_rows": len(sub["rows"])}


@router.post("/solutions-store/factor-submissions/{submission_id}/reject")
async def reject_factor_submission(submission_id: str, request: Request, user: dict = Depends(get_current_user)):
    if not _is_admin(user):
        raise HTTPException(status_code=403, detail="Admin approval required")
    body = {}
    try:
        body = await request.json()
    except Exception:
        pass
    sub = await db.store_factor_submissions.find_one({"submission_id": submission_id})
    if not sub:
        raise HTTPException(status_code=404, detail="Submission not found")
    await db.store_factor_submissions.update_one(
        {"submission_id": submission_id},
        {"$set": {"status": "rejected", "reviewed_by": user["user_id"], "reviewed_at": _now(),
                  "note": (body.get("note") or "").strip() or None}},
    )
    return {"status": "rejected"}
