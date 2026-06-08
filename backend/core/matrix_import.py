"""Comparison-MATRIX import helpers (Step-2 'Import from XLS / Google Sheet').

Expected layout (same shape the URL importer produces):
    Row 1  : header  → [Option, <Factor 1>, <Factor 2>, ...]
    Row 2+ : data    → [<Option name>, <value>, <value>, ...]

Resilient by design:
  * header with factor names but NO data rows  → import FACTORS only
  * data rows but empty value cells            → import factors + OPTIONS (no values)
"""
from __future__ import annotations

import io
import re
import csv
from typing import Any, Dict, List, Optional, Tuple

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

_OPTION_HEADERS = {"option", "options", "name", "item", "product", "scheme", "fund", "candidate"}


def extract_rows_from_xlsx(data: bytes) -> List[List[Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def extract_rows_from_csv(data: bytes) -> List[List[Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [list(r) for r in csv.reader(io.StringIO(text))]


def parse_matrix(rows: List[List[Any]]) -> Dict[str, Any]:
    """Parse raw rows → {"factor_names": [...], "candidates": [{name, attributes}]}.
    Raises ValueError when there is nothing usable."""
    # Drop fully-empty rows.
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("The sheet is empty.")

    header = [str(c).strip() for c in rows[0]]
    # Factor names = header columns after the first (option) column.
    factor_names: List[str] = []
    seen: set = set()
    for h in header[1:]:
        h = h.strip()
        if not h:
            continue
        key = h.lower()
        if key in seen:
            continue
        seen.add(key)
        factor_names.append(h)

    if not factor_names:
        raise ValueError("No factor columns found. Put factor names in row 1 from column B onwards.")

    candidates: List[Dict[str, Any]] = []
    opt_seen: set = set()
    for r in rows[1:]:
        name = str(r[0]).strip() if r else ""
        if not name or name.lower() in opt_seen:
            continue
        opt_seen.add(name.lower())
        attrs: Dict[str, Any] = {}
        for i, fn in enumerate(factor_names):
            ci = i + 1  # +1 to skip the option column
            val = str(r[ci]).strip() if ci < len(r) else ""
            if val:
                attrs[fn] = val
        candidates.append({"name": name[:120], "attributes": attrs})

    return {"factor_names": factor_names, "candidates": candidates}


# ── Google Sheets (public link → CSV export) ────────────────────────────────
_SHEET_ID_RE = re.compile(r"/spreadsheets/d/([a-zA-Z0-9-_]+)")
_GID_RE = re.compile(r"[#&?]gid=(\d+)")


def gsheet_id_and_gid(url: str) -> Tuple[Optional[str], Optional[str]]:
    m = _SHEET_ID_RE.search(url or "")
    g = _GID_RE.search(url or "")
    return (m.group(1) if m else None, g.group(1) if g else None)


async def fetch_public_gsheet_rows(url: str) -> List[List[Any]]:
    """Read a link-shared (public) Google Sheet as CSV — no OAuth needed."""
    sid, gid = gsheet_id_and_gid(url)
    if not sid:
        raise ValueError("That doesn't look like a Google Sheets link.")
    export = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=csv"
    if gid:
        export += f"&gid={gid}"
    async with httpx.AsyncClient(timeout=25.0, follow_redirects=True) as cli:
        r = await cli.get(export)
    ctype = r.headers.get("content-type", "")
    if r.status_code != 200 or "text/csv" not in ctype:
        # Not public (login wall) → signal caller to try OAuth.
        raise PermissionError("Sheet is not link-shared publicly.")
    return extract_rows_from_csv(r.content)


# ── Downloadable template (.xlsx) ───────────────────────────────────────────
def build_template_xlsx(factor_names: List[str], option_names: List[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Decision Matrix"

    factors = factor_names or ["Factor 1", "Factor 2", "Factor 3"]
    options = option_names or ["Option A", "Option B", "Option C"]
    header = ["Option"] + factors
    ws.append(header)

    head_fill = PatternFill("solid", fgColor="1E293B")
    head_font = Font(bold=True, color="FFFFFF")
    for col, _ in enumerate(header, start=1):
        c = ws.cell(row=1, column=col)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center")

    for nm in options:
        ws.append([nm] + ["" for _ in factors])

    ws.column_dimensions["A"].width = 26
    for i in range(len(factors)):
        ws.column_dimensions[chr(ord("B") + i)].width = 18
    ws.freeze_panes = "B2"

    note = ws.cell(row=len(options) + 3, column=1,
                   value="How to fill: Column A = your options (rows). Row 1 (B onwards) = factor "
                         "names. Cells = each option's value for that factor (numbers like 24.5, "
                         "₹899, 4.2 are auto-scored). Options or values may be left blank.")
    note.font = Font(italic=True, color="64748B")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
