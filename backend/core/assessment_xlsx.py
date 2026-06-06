"""Shared XLSX template builder + parser for assessment data collection.

Used by both the My Dezider and Pros & Cons flows (Phase C). The template
captures, per option, an "Actual Value" and an "Assess %" for every factor and
sub-factor. Factors and sub-factors are visually differentiated (main factor
rows are bold + shaded; sub-factor rows are indented with a "↳" prefix and a
"Sub-factor" level tag).

Robust round-trip: row 1 is a HIDDEN machine-header that encodes the column
contract (factor-id column + per-option field/option-id), so import does not
depend on column order or human header text. Row 2 is the human header.
Data rows start at row 3.
"""

from io import BytesIO
from typing import Any, Callable, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


# Machine-header tokens (row 1, hidden)
_KEY_FACTOR_ID = "__factor_id__"
_KEY_ACTUAL = "actual:"   # + option_id
_KEY_PCT = "pct:"         # + option_id

_MAIN_FILL = PatternFill("solid", fgColor="E8EEF7")
_HEAD_FILL = PatternFill("solid", fgColor="1F2A44")
_LOCK_FILL = PatternFill("solid", fgColor="F3F4F6")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
_MAIN_FONT = Font(bold=True, color="1F2A44")
_SUB_FONT = Font(color="334155")
_THIN = Side(style="thin", color="D0D7E2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_template(
    title: str,
    factors: List[Dict[str, Any]],
    options: List[Dict[str, Any]],
    get_cell: Callable[[str, str], Dict[str, Any]],
) -> bytes:
    """Build an .xlsx assessment template.

    factors: ordered list of dicts with keys:
        id, name, parent_id (None for main), expected, unit
        (caller MUST pass them already ordered: each main factor immediately
         followed by its sub-factors).
    options: list of dicts with keys: id, name.
    get_cell(option_id, factor_id) -> {"actual": str, "pct": int} (existing values).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Assessment"

    # ── Column layout ────────────────────────────────────────────────
    # A: Factor ID (locked key) | B: Level | C: Factor | D: Parent
    # E: Expected | F: Unit | then per option: Actual, Assess %
    fixed_machine = [_KEY_FACTOR_ID, "level", "factor", "parent", "expected", "unit"]
    fixed_human = ["Factor ID (do not edit)", "Level", "Factor / Sub-factor", "Parent Factor", "Expected / Target", "Unit"]

    machine_row: List[str] = list(fixed_machine)
    human_row: List[str] = list(fixed_human)
    for opt in options:
        machine_row.append(f"{_KEY_ACTUAL}{opt['id']}")
        human_row.append(f"{opt['name']} — Actual Value")
        machine_row.append(f"{_KEY_PCT}{opt['id']}")
        human_row.append(f"{opt['name']} — Assess %")

    # Row 1 — machine header (hidden)
    ws.append(machine_row)
    # Row 2 — human header (styled)
    ws.append(human_row)
    for c in range(1, len(human_row) + 1):
        cell = ws.cell(row=2, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER

    # Optional 0-100 validation for the Assess % columns
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
    ws.add_data_validation(dv)

    # ── Data rows ────────────────────────────────────────────────────
    r = 3
    pct_col_letters: List[str] = []
    for f in factors:
        is_sub = bool(f.get("parent_id"))
        name = f.get("name") or ""
        ws.cell(row=r, column=1, value=f.get("id"))
        ws.cell(row=r, column=2, value="Sub-factor" if is_sub else "Main Factor")
        ws.cell(row=r, column=3, value=(f"    \u21b3 {name}" if is_sub else name))
        ws.cell(row=r, column=4, value=f.get("parent_name") or "")
        ws.cell(row=r, column=5, value=f.get("expected") if f.get("expected") not in (None, "") else "")
        ws.cell(row=r, column=6, value=f.get("unit") or "")

        col = 7
        for opt in options:
            existing = get_cell(opt["id"], f["id"]) or {}
            a_cell = ws.cell(row=r, column=col, value=existing.get("actual") or "")
            a_cell.border = _BORDER
            p_cell = ws.cell(row=r, column=col + 1)
            pv = existing.get("pct")
            if pv not in (None, "", 0):
                p_cell.value = int(pv)
            p_cell.border = _BORDER
            dv.add(p_cell)
            if r == 3:
                pct_col_letters.append(p_cell.column_letter)
            col += 2

        # Row styling — differentiate main vs sub
        for c in range(1, len(machine_row) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            if c <= 6:
                cell.font = _MAIN_FONT if not is_sub else _SUB_FONT
            if not is_sub:
                cell.fill = _MAIN_FILL
            if c == 1:  # lock-ish look for the key column
                cell.fill = _LOCK_FILL
        r += 1

    # Row 1 hidden + freeze panes + widths
    ws.row_dimensions[1].hidden = True
    ws.freeze_panes = "C3"
    widths = [20, 12, 34, 20, 18, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
    for c in range(7, len(machine_row) + 1):
        ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = 18

    # Title note in a second sheet (kept minimal)
    info = wb.create_sheet("Read me")
    info["A1"] = f"Assessment template — {title}"
    info["A1"].font = Font(bold=True, size=13)
    info["A3"] = "• Fill the 'Actual Value' and 'Assess %' (0-100) columns per option."
    info["A4"] = "• Assess % is optional — leave blank to skip; sub-factors roll up into their main factor."
    info["A5"] = "• Do NOT edit the 'Factor ID' column or the 'Assessment' sheet headers."
    info.column_dimensions["A"].width = 90

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_value_matrix(
    factors: List[Dict[str, Any]],
    options: List[Dict[str, Any]],
    get_cell: Callable[[str, str], Dict[str, Any]],
) -> List[List[Any]]:
    """Return the assessment template as a plain 2D matrix (list of rows),
    using the SAME column contract as build_template():
      row 0 = machine header (hidden on render), row 1 = human header,
      rows 2.. = data. Used by the Google Sheets writer so XLS and Google
      Sheets share one import parser (parse_rows)."""
    fixed_machine = [_KEY_FACTOR_ID, "level", "factor", "parent", "expected", "unit"]
    fixed_human = ["Factor ID (do not edit)", "Level", "Factor / Sub-factor", "Parent Factor", "Expected / Target", "Unit"]

    machine_row: List[Any] = list(fixed_machine)
    human_row: List[Any] = list(fixed_human)
    for opt in options:
        machine_row.append(f"{_KEY_ACTUAL}{opt['id']}")
        human_row.append(f"{opt.get('name') or 'Option'} — Actual Value")
        machine_row.append(f"{_KEY_PCT}{opt['id']}")
        human_row.append(f"{opt.get('name') or 'Option'} — Assess %")

    matrix: List[List[Any]] = [machine_row, human_row]
    for f in factors:
        is_sub = bool(f.get("parent_id"))
        name = f.get("name") or ""
        row: List[Any] = [
            f.get("id"),
            "Sub-factor" if is_sub else "Main Factor",
            (f"    \u21b3 {name}" if is_sub else name),
            f.get("parent_name") or "",
            f.get("expected") if f.get("expected") not in (None, "") else "",
            f.get("unit") or "",
        ]
        for opt in options:
            existing = get_cell(opt["id"], f["id"]) or {}
            row.append(existing.get("actual") or "")
            pv = existing.get("pct")
            row.append(int(pv) if pv not in (None, "", 0) else "")
        matrix.append(row)
    return matrix


def parse_rows(rows: List[List[Any]]) -> List[Dict[str, Any]]:
    """Parse an assessment matrix (list of rows; row 0 = machine header) into
    a list of {factor_id, option_id, actual, assessment_pct}. Shared by the
    XLSX parser and the Google Sheets importer."""
    if not rows or len(rows) < 3:
        return []

    machine = [str(c) if c is not None else "" for c in rows[0]]
    factor_id_col: Optional[int] = None
    col_map: Dict[int, Dict[str, str]] = {}
    for idx, token in enumerate(machine):
        if token == _KEY_FACTOR_ID:
            factor_id_col = idx
        elif token.startswith(_KEY_ACTUAL):
            col_map[idx] = {"field": "actual", "option_id": token[len(_KEY_ACTUAL):]}
        elif token.startswith(_KEY_PCT):
            col_map[idx] = {"field": "pct", "option_id": token[len(_KEY_PCT):]}

    if factor_id_col is None:
        raise ValueError("Invalid template: missing factor-id header. Re-download the template.")

    acc: Dict[str, Dict[str, Any]] = {}
    for row in rows[2:]:
        if factor_id_col >= len(row):
            continue
        fid = row[factor_id_col]
        if not fid:
            continue
        fid = str(fid).strip()
        for col_idx, meta in col_map.items():
            if col_idx >= len(row):
                continue
            val = row[col_idx]
            if val is None or str(val).strip() == "":
                continue
            key = f"{meta['option_id']}::{fid}"
            entry = acc.setdefault(key, {"factor_id": fid, "option_id": meta["option_id"]})
            if meta["field"] == "actual":
                entry["actual"] = str(val).strip()
            else:
                try:
                    entry["assessment_pct"] = max(0, min(100, int(round(float(val)))))
                except (ValueError, TypeError):
                    pass
    return list(acc.values())


def parse_template(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Parse a filled .xlsx template → list of
    {factor_id, option_id, actual, assessment_pct}. Robust to column reordering
    (uses the hidden machine-header in row 1)."""
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    if "Assessment" in wb.sheetnames:
        ws = wb["Assessment"]
    else:
        ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    return parse_rows([list(r) for r in rows])
