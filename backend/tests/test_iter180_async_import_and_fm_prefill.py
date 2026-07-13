"""ITER 180 — Async import job + FM template prefill regression suite.

Covers:
  * File Import (Step 2) — async job start/poll happy path + error cases.
  * File Import OLD sync endpoint regression — still returns factors/options.
  * Financial Model template prefill — ?model_id= pre-fills assumption values.
  * Financial Model import-file + import-sheet regression.

Auth: super@test.com / SuperPass2026!  → POST /api/auth/login → session_token.
Chunked upload payload key: chunk_b64.
"""
from __future__ import annotations

import base64
import io
import os
import time

import pytest
import requests

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL", "https://repo-blueprint-1.preview.emergentagent.com").rstrip("/")
EMAIL = "super@test.com"
PASSWORD = "SuperPass2026!"


# ─── Auth + helpers ──────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def token() -> str:
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": EMAIL, "password": PASSWORD}, timeout=20)
    assert r.status_code == 200, r.text
    tok = r.json().get("session_token") or r.json().get("token")
    assert tok, f"No session_token in login response: {r.text[:200]}"
    return tok


@pytest.fixture(scope="module")
def auth(token):
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {token}", "Content-Type": "application/json"})
    return s


def _chunk_upload(auth, raw: bytes, filename: str) -> str:
    init = auth.post(f"{BASE_URL}/api/uploads/init",
                     json={"filename": filename, "size": len(raw)}, timeout=20)
    assert init.status_code == 200, init.text
    upload_id = init.json()["upload_id"]
    b64 = base64.b64encode(raw).decode("ascii")
    ch = auth.post(f"{BASE_URL}/api/uploads/chunk",
                   json={"upload_id": upload_id, "index": 0, "chunk_b64": b64, "final": True},
                   timeout=30)
    assert ch.status_code == 200, ch.text
    return upload_id


@pytest.fixture(scope="module")
def decision_id(auth) -> str:
    # Reuse an existing decision for super@test.com (request notes one exists).
    r = auth.get(f"{BASE_URL}/api/decisions", timeout=20)
    assert r.status_code == 200, r.text
    body = r.json()
    rows = body if isinstance(body, list) else (body.get("decisions") or body.get("items") or [])
    if rows:
        return rows[0]["id"]
    # Fallback: create one.
    c = auth.post(f"{BASE_URL}/api/decisions",
                  json={"title": "TEST_iter180_async_import",
                        "description": "iter180 async import probe"}, timeout=20)
    assert c.status_code in (200, 201), c.text
    return c.json()["id"] if isinstance(c.json(), dict) else c.json()


# ─── 1. Async import job — happy path (small TXT) ───────────────────────────
class TestAsyncImportJob:
    """POST /file-import/decision/{id}/start → GET /file-import/jobs/{job_id}"""

    TXT_BODY = (
        "VC Firm Comparison\n"
        "We are picking a seed-stage investor.\n\n"
        "Sequoia Capital India: stage Seed/A, sector Tech, ticket $1-5M.\n"
        "Accel Partners: stage Seed, sector SaaS, ticket $0.5-3M.\n"
        "Lightspeed Venture Partners: stage A/B, sector Consumer, ticket $5-10M.\n"
        "\nFactors that matter: stage focus, sector focus, ticket size, fund size, value-add.\n"
    ).encode("utf-8")

    def test_start_and_poll_to_done(self, auth, decision_id):
        upload_id = _chunk_upload(auth, self.TXT_BODY, "vc_firms.txt")
        start = auth.post(
            f"{BASE_URL}/api/file-import/decision/{decision_id}/start",
            json={"filename": "vc_firms.txt", "upload_id": upload_id,
                  "ai_tier": "fast", "crawl_web": False}, timeout=20)
        assert start.status_code == 200, start.text
        job_id = start.json().get("job_id")
        assert job_id and isinstance(job_id, str)

        # Poll up to ~90s.
        last = None
        for _ in range(45):
            time.sleep(2)
            g = auth.get(f"{BASE_URL}/api/file-import/jobs/{job_id}", timeout=15)
            assert g.status_code == 200, g.text
            last = g.json()
            assert "status" in last and "stage" in last
            if last["status"] in ("done", "error"):
                break
        assert last is not None
        assert last["status"] == "done", f"Job not done: {last}"
        res = last.get("result") or {}
        assert isinstance(res.get("factors"), list)
        assert isinstance(res.get("options"), list)
        assert (len(res["factors"]) + len(res["options"])) > 0


# ─── 2. Async start/get error codes ─────────────────────────────────────────
class TestAsyncImportErrors:
    def test_get_nonexistent_job_returns_404(self, auth):
        r = auth.get(f"{BASE_URL}/api/file-import/jobs/does-not-exist-xyz", timeout=15)
        assert r.status_code == 404, r.text

    def test_start_missing_upload_id_returns_400(self, auth, decision_id):
        r = auth.post(
            f"{BASE_URL}/api/file-import/decision/{decision_id}/start",
            json={"filename": "x.txt", "ai_tier": "fast", "crawl_web": False}, timeout=15)
        assert r.status_code == 400, r.text

    def test_start_nonexistent_decision_returns_404(self, auth):
        # Pretend chunked upload exists; backend should 404 on decision lookup first.
        r = auth.post(
            f"{BASE_URL}/api/file-import/decision/does-not-exist-decision/start",
            json={"filename": "x.txt", "upload_id": "nope", "ai_tier": "fast"}, timeout=15)
        assert r.status_code == 404, r.text


# ─── 3. OLD sync endpoint regression ────────────────────────────────────────
class TestSyncImportRegression:
    TXT_BODY = (
        "Choosing a CRM\n"
        "Salesforce: enterprise SaaS, strong ecosystem.\n"
        "HubSpot: SMB-friendly, free tier.\n"
        "Zoho CRM: low-cost, all-in-one.\n"
        "Factors: price, ease of use, integrations, support.\n"
    ).encode("utf-8")

    def test_old_sync_endpoint_still_works(self, auth, decision_id):
        upload_id = _chunk_upload(auth, self.TXT_BODY, "crm_compare.txt")
        r = auth.post(
            f"{BASE_URL}/api/file-import/decision/{decision_id}",
            json={"filename": "crm_compare.txt", "upload_id": upload_id,
                  "ai_tier": "fast", "crawl_web": False}, timeout=120)
        assert r.status_code == 200, r.text
        body = r.json()
        assert isinstance(body.get("factors"), list)
        assert isinstance(body.get("options"), list)
        assert (len(body["factors"]) + len(body["options"])) > 0


# ─── 4. FM template prefill via ?model_id=  ─────────────────────────────────
class TestTemplatePrefill:
    @pytest.fixture(scope="class")
    def model_id(self, auth):
        """Create a financial model on the pre-seeded Org with custom assumptions."""
        # Reuse seeded org for super@test.com (per iter179 context).
        SEED_ORG = "c7941866-ca40-4a68-af9f-c232cd2c6c32"
        SEED_GOAL = "dd23b482-9109-4878-a53b-5e24e0af7474"
        # Sanity: ensure the org is visible to this account.
        org_list = auth.get(f"{BASE_URL}/api/seven-seven/orgs", timeout=15).json()
        org_ids = [o["id"] for o in (org_list.get("orgs") or [])]
        if SEED_ORG not in org_ids:
            if not org_ids:
                # Bootstrap a new org as a fallback.
                org = auth.post(f"{BASE_URL}/api/seven-seven/orgs",
                                json={"name": "TEST_iter180_org",
                                      "org_type": "BUSINESS",
                                      "life_area": "business"}, timeout=20)
                assert org.status_code in (200, 201), org.text
                org_id = org.json()["user_org"]["id"]
            else:
                org_id = org_ids[0]
        else:
            org_id = SEED_ORG

        assumptions = {
            "year1_revenue": 7654321,
            "gross_margin_pct": 42,
            "opex_pct": 27,
            "wacc_pct": 19,
            "tax_rate_pct": 25,
            "beta": 1.4,
            "risk_free_pct": 7.25,
        }
        m = auth.post(f"{BASE_URL}/api/financial-models",
                      json={"user_org_id": org_id, "leg_goal_id": SEED_GOAL,
                            "name": "TEST_iter180_model", "assumptions": assumptions}, timeout=20)
        assert m.status_code in (200, 201), m.text
        body = m.json()
        return body.get("id") or body.get("model", {}).get("id")

    def _xlsx_value_for_label(self, raw: bytes, label_substr: str):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), data_only=False)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if not row:
                continue
            cell0 = str(row[0] or "")
            if label_substr.lower() in cell0.lower():
                return row[1] if len(row) > 1 else None
        return None

    def test_blank_template_without_model_id(self, auth):
        r = auth.get(f"{BASE_URL}/api/financial-models/templates/inputs.xlsx", timeout=20)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("content-type", "")
        # Blank template — value cell for an arbitrary label should not match our
        # custom probe value (7,654,321).
        v = self._xlsx_value_for_label(r.content, "Year-1 revenue")
        assert v != 7654321

    def test_template_with_model_id_prefills_assumption_values(self, auth, model_id):
        r = auth.get(
            f"{BASE_URL}/api/financial-models/templates/inputs.xlsx?model_id={model_id}",
            timeout=20)
        assert r.status_code == 200, r.text
        v_rev = self._xlsx_value_for_label(r.content, "Year-1 revenue")
        v_gm = self._xlsx_value_for_label(r.content, "Gross margin")
        v_wacc = self._xlsx_value_for_label(r.content, "WACC")
        # If template-prefill is wired correctly, the saved values must appear.
        assert v_rev == 7654321, (
            f"Expected 7654321 (saved assumption) in Year-1 revenue, got {v_rev!r}. "
            "Likely root-cause: download_template() looks up model by 'owner_user_id' "
            "but documents are stored with 'user_id'.")
        assert v_gm == 42, f"Expected 42, got {v_gm!r}"
        assert v_wacc == 19, f"Expected 19, got {v_wacc!r}"


# ─── 5. FM import-file + import-sheet regression ────────────────────────────
class TestFinModelImportsRegression:
    def test_import_file_round_trip(self, auth):
        # Build a minimal filled template in memory and upload it.
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Label", "Value"])
        ws.append(["Year-1 revenue", 9000000])
        ws.append(["Gross margin %", 55])
        ws.append(["WACC %", 18])
        buf = io.BytesIO()
        wb.save(buf)
        raw = buf.getvalue()
        upload_id = _chunk_upload(auth, raw, "filled_template.xlsx")
        r = auth.post(f"{BASE_URL}/api/financial-models/import-file",
                      json={"filename": "filled_template.xlsx", "upload_id": upload_id}, timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["patch"].get("year1_revenue") == 9000000
        assert body["patch"].get("gross_margin_pct") == 55
        assert body["patch"].get("wacc_pct") == 18

    def test_import_sheet_bad_url_returns_422(self, auth):
        r = auth.post(f"{BASE_URL}/api/financial-models/import-sheet",
                      json={"sheet_url": "https://example.com/not-a-sheet"}, timeout=20)
        assert r.status_code == 422, r.text
