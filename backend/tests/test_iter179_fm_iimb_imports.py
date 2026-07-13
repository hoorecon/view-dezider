"""ITER 179 backend tests — Financial Model: Template, Import (file/sheet), IIMB.

Covers:
  • GET  /api/financial-models/templates/inputs.xlsx           (200, xlsx mime, attachment, parsable, contains template labels)
  • POST /api/financial-models/import-file                     (parses a filled template via chunked upload; 422 on non-template)
  • POST /api/financial-models/import-sheet                    (422 on bad URL; valid docs.google.com/spreadsheets/... URL is accepted by the id parser — auth gate returns >=400 but not 422 'paste a valid...')
  • POST /api/financial-models/compute (iimb block, CAPM vs direct WACC)
"""
import io
import os
import uuid
from typing import Any, Dict

import pytest
import requests

BASE_URL = (os.environ.get("EXPO_PUBLIC_BACKEND_URL")
            or "https://repo-blueprint-1.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

OWNER = {"email": "super@test.com", "password": "SuperPass2026!"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ───────────────────────────── Fixtures ─────────────────────────────
@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    r = s.post(f"{API}/auth/login", json=OWNER, timeout=30)
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    j = r.json()
    tok = j.get("access_token") or j.get("token") or j.get("session_token")
    assert tok, f"No token in login: {j}"
    s.headers["Authorization"] = f"Bearer {tok}"
    return s


# Helper: upload bytes via /api/uploads/init + /api/uploads/chunk → upload_id
def _chunked_upload(session, filename: str, content: bytes) -> str:
    import base64
    init = session.post(f"{API}/uploads/init",
                        json={"filename": filename, "total_size": len(content),
                              "mime": XLSX_MIME if filename.endswith(".xlsx") else "text/plain"},
                        timeout=30)
    assert init.status_code == 200, f"upload init failed: {init.status_code} {init.text}"
    upload_id = init.json().get("upload_id") or init.json().get("id")
    assert upload_id, f"no upload_id: {init.json()}"
    # single chunk
    b64 = base64.b64encode(content).decode("ascii")
    ch = session.post(f"{API}/uploads/chunk",
                      json={"upload_id": upload_id, "index": 0, "chunk_b64": b64, "is_last": True},
                      timeout=60)
    assert ch.status_code == 200, f"upload chunk failed: {ch.status_code} {ch.text}"
    return upload_id


# ─────────────── 1. Template download ───────────────
class TestTemplateDownload:
    def test_template_xlsx_200_and_mime(self, session):
        r = session.get(f"{API}/financial-models/templates/inputs.xlsx", timeout=30)
        assert r.status_code == 200, f"template dl: {r.status_code} {r.text[:200]}"
        assert XLSX_MIME in r.headers.get("content-type", ""), r.headers
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower() and "financial-model-template.xlsx" in cd, cd
        assert len(r.content) > 1000, f"unexpectedly small: {len(r.content)} bytes"
        # parsable as xlsx with labelled rows in column A
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        labels_col_a = [(ws.cell(row=r_, column=1).value or "") for r_ in range(1, ws.max_row + 1)]
        joined = " | ".join(str(x) for x in labels_col_a if x)
        # spot-check a few core labels
        for needle in ["Year-1 revenue", "Gross margin", "WACC", "Beta", "Market risk premium",
                       "Cost of debt", "Risk-free rate"]:
            assert needle in joined, f"label missing in template: {needle}"


# ─────────────── 2. Import file (good path + 422) ───────────────
class TestImportFile:
    def test_import_filled_template_returns_patch(self, session):
        # Build a filled copy of the downloaded template
        r = session.get(f"{API}/financial-models/templates/inputs.xlsx", timeout=30)
        assert r.status_code == 200
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # write a few values: scalar Year-1 revenue (col B) and a 5-year revenue_by_year series
        for row in range(1, ws.max_row + 1):
            lbl = (ws.cell(row=row, column=1).value or "")
            if lbl == "Year-1 revenue":
                ws.cell(row=row, column=2, value=12_000_000)
            elif lbl == "Revenue by year (Y1..Y5)":
                for c, v in enumerate([10_000_000, 15_000_000, 22_000_000, 30_000_000, 40_000_000], 2):
                    ws.cell(row=row, column=c, value=v)
            elif lbl == "Gross margin %":
                ws.cell(row=row, column=2, value=55)
            elif lbl == "WACC %":
                ws.cell(row=row, column=2, value=18)
        buf = io.BytesIO()
        wb.save(buf)
        filled = buf.getvalue()
        filename = f"TEST_iter179_filled_{uuid.uuid4().hex[:6]}.xlsx"
        uid = _chunked_upload(session, filename, filled)
        resp = session.post(f"{API}/financial-models/import-file",
                            json={"filename": filename, "upload_id": uid}, timeout=60)
        assert resp.status_code == 200, f"import-file: {resp.status_code} {resp.text}"
        j = resp.json()
        assert "patch" in j and "found" in j, j
        patch: Dict[str, Any] = j["patch"]
        # scalar pulled
        assert patch.get("year1_revenue") == 12_000_000
        assert patch.get("gross_margin_pct") == 55
        assert patch.get("wacc_pct") == 18
        # series pulled (length 5)
        rev = patch.get("revenue_by_year")
        assert isinstance(rev, list) and len(rev) == 5 and rev[0] == 10_000_000 and rev[-1] == 40_000_000

    def test_import_non_template_returns_422(self, session):
        # An XLSX with garbage labels in col A → parser finds zero matches → 422
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws["A1"] = "Foo"; ws["B1"] = 1
        ws["A2"] = "Bar"; ws["B2"] = 2
        buf = io.BytesIO(); wb.save(buf)
        filename = f"TEST_iter179_garbage_{uuid.uuid4().hex[:6]}.xlsx"
        uid = _chunked_upload(session, filename, buf.getvalue())
        resp = session.post(f"{API}/financial-models/import-file",
                            json={"filename": filename, "upload_id": uid}, timeout=30)
        assert resp.status_code == 422, f"expected 422 on non-template: {resp.status_code} {resp.text}"


# ─────────────── 3. Import sheet ───────────────
class TestImportSheet:
    def test_import_sheet_bad_url_422(self, session):
        resp = session.post(f"{API}/financial-models/import-sheet",
                            json={"sheet_url": "https://example.com/not-a-sheet"}, timeout=30)
        assert resp.status_code == 422, f"expected 422 on bad url: {resp.status_code} {resp.text}"
        assert "valid google sheets" in resp.text.lower()

    def test_import_sheet_empty_url_422(self, session):
        resp = session.post(f"{API}/financial-models/import-sheet",
                            json={"sheet_url": ""}, timeout=30)
        assert resp.status_code == 422, f"expected 422: {resp.status_code} {resp.text}"

    def test_import_sheet_valid_id_url_is_accepted_by_id_parser(self, session):
        # We don't have a real public sheet; but a *valid-shaped* docs URL must NOT 422 with
        # "Paste a valid Google Sheets link." It should proceed (and then 4xx for fetch failure
        # such as 422 with a fetch error msg OR 403 require connect google). The key is it must
        # NOT be rejected by the id parser ("Paste a valid Google Sheets link.").
        url = "https://docs.google.com/spreadsheets/d/1aBcDeFGhiJklmNopQrsTuVwxYz0123456789AbCd/edit#gid=0"
        resp = session.post(f"{API}/financial-models/import-sheet",
                            json={"sheet_url": url}, timeout=60)
        # Either 422 (fetch failure with different msg), 403 (private+no creds), or 200 — but NOT
        # the "paste a valid" 422 from the id parser. Accept any non-id-parser response.
        body = resp.text.lower()
        assert "paste a valid google sheets" not in body, \
            f"id parser incorrectly rejected a well-formed URL: {resp.status_code} {resp.text}"


# ─────────────── 4. Compute IIMB block ───────────────
class TestComputeIIMB:
    _ASSUMPT_BASE = {
        "revenue_by_year": [10_000_000, 18_000_000, 30_000_000, 46_000_000, 65_000_000],
        "gross_margin_pct": 55, "opex_pct": 35, "tax_rate_pct": 25,
        "opening_gross_block": 3_000_000, "capex_by_year": [2_000_000] * 5,
        "depreciation_pct": 15,
        "debtor_days": 45, "inventory_days": 30, "creditor_days": 40,
        "opening_debt": 2_000_000, "interest_rate_pct": 12,
        "opening_equity_capital": 8_000_000, "opening_cash": 7_000_000,
        "shares_outstanding": 800_000,
        # WACC build-up
        "risk_free_pct": 7, "beta": 1.2, "market_risk_premium_pct": 6,
        "cost_of_debt_pct": 10, "market_cap": 50_000_000,
        "minority_interest": 1_000_000, "preference_capital": 500_000,
        "non_operating_assets": 200_000,
        "wacc_pct": 18, "terminal_growth_pct": 4,
    }

    def _compute(self, session, assumptions, years=5):
        r = session.post(f"{API}/financial-models/compute",
                         json={"assumptions": assumptions, "projection_years": years}, timeout=30)
        assert r.status_code == 200, f"compute: {r.status_code} {r.text}"
        return r.json()["computed"]

    def test_iimb_block_shape(self, session):
        a = dict(self._ASSUMPT_BASE); a["wacc_mode"] = "capm"
        c = self._compute(session, a)
        assert "iimb" in c, c.keys()
        iimb = c["iimb"]
        for k in ["cost_of_equity_pct", "equity_weight_pct", "debt_weight_pct",
                  "wacc_capm_pct", "wacc_used_pct",
                  "noplat", "gross_cash_flow", "increase_in_nwc", "fcff", "pv_fcff",
                  "enterprise_value", "equity_value", "per_share",
                  "simple_market_ev", "fuller_market_ev"]:
            assert k in iimb, f"iimb missing key {k}"
        for k in ["noplat", "gross_cash_flow", "increase_in_nwc", "fcff", "pv_fcff"]:
            assert isinstance(iimb[k], list) and len(iimb[k]) == 5, f"{k} should be 5-list, got {iimb[k]}"

    def test_wacc_capm_math_and_used_equals_capm(self, session):
        a = dict(self._ASSUMPT_BASE); a["wacc_mode"] = "capm"
        c = self._compute(session, a)
        iimb = c["iimb"]
        # Cost of equity = rf + beta * mrp = 7 + 1.2*6 = 14.2
        assert abs(iimb["cost_of_equity_pct"] - 14.2) < 0.01, iimb["cost_of_equity_pct"]
        # weights from mkt_cap=50M, open_debt=2M → we=50/52≈96.15, wd≈3.85
        assert abs(iimb["equity_weight_pct"] - 96.15) < 0.1, iimb["equity_weight_pct"]
        assert abs(iimb["debt_weight_pct"] - 3.85) < 0.1, iimb["debt_weight_pct"]
        # kd after tax = 10 * (1-0.25) = 7.5 → wacc_capm = 0.9615*14.2 + 0.0385*7.5 ≈ 13.94
        expected_wacc_capm = 0.9615 * 14.2 + 0.0385 * 7.5
        assert abs(iimb["wacc_capm_pct"] - expected_wacc_capm) < 0.1, \
            f"{iimb['wacc_capm_pct']} vs expected {expected_wacc_capm}"
        # In capm mode, wacc_used should equal wacc_capm
        assert abs(iimb["wacc_used_pct"] - iimb["wacc_capm_pct"]) < 0.01, \
            f"wacc_used {iimb['wacc_used_pct']} != wacc_capm {iimb['wacc_capm_pct']}"
        # And the valuation block's wacc_pct should also reflect CAPM
        assert abs(c["valuation"]["wacc_pct"] - iimb["wacc_capm_pct"]) < 0.01

    def test_wacc_direct_used_equals_wacc_pct(self, session):
        a = dict(self._ASSUMPT_BASE); a["wacc_mode"] = "direct"; a["wacc_pct"] = 18
        c = self._compute(session, a)
        iimb = c["iimb"]
        # In direct mode wacc_used ≈ 18
        assert abs(iimb["wacc_used_pct"] - 18) < 0.01, iimb["wacc_used_pct"]
        # wacc_capm is still reported but should NOT equal wacc_used here
        # (unless coincidentally equal — which is highly unlikely with these inputs)
        assert abs(iimb["wacc_used_pct"] - iimb["wacc_capm_pct"]) > 0.5, \
            f"capm vs used unexpectedly equal: used={iimb['wacc_used_pct']} capm={iimb['wacc_capm_pct']}"

    def test_fcff_matches_valuation_block(self, session):
        a = dict(self._ASSUMPT_BASE); a["wacc_mode"] = "capm"
        c = self._compute(session, a)
        # iimb.fcff and valuation.fcff should be identical lists
        assert c["iimb"]["fcff"] == c["valuation"]["fcff"], (c["iimb"]["fcff"], c["valuation"]["fcff"])

    def test_market_ev_crosscheck(self, session):
        a = dict(self._ASSUMPT_BASE); a["wacc_mode"] = "capm"
        c = self._compute(session, a)
        iimb = c["iimb"]
        # simple_market_ev = mkt_cap + (open_debt - open_cash) = 50M + (2M - 7M) = 45M
        assert abs(iimb["simple_market_ev"] - 45_000_000) < 1, iimb["simple_market_ev"]
        # fuller_market_ev = mkt_cap + open_debt + minority + pref - cash - non_op
        #                  = 50M + 2M + 1M + 0.5M - 7M - 0.2M = 46.3M
        assert abs(iimb["fuller_market_ev"] - 46_300_000) < 1, iimb["fuller_market_ev"]
